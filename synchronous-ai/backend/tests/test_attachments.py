"""Task attachments: upload validation, linking, access control, and what the agent receives."""

from __future__ import annotations

import io
from datetime import timedelta
from types import SimpleNamespace

from sca.db import session_scope
from sca.models import Agent, Attachment, utcnow
from sca.services.tasks import submit_task
from tests.conftest import PASSWORD, login_as, make_agent


def _png(w: int = 4, h: int = 4) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (w, h), (200, 30, 30)).save(buf, format="PNG")
    return buf.getvalue()


def _docx(text: str) -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["region", "revenue"])
    ws.append(["EMEA", 1200])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(api, name: str, data: bytes):
    return api.post("/api/attachments", files={"file": (name, data)})


def test_upload_validation_link_and_access(admin, server, offline_provider, monkeypatch):
    # The task submitted below only checks linking; keep workers from running it (offline provider).
    from sca.orchestrator import executor

    monkeypatch.setattr(executor, "execute_task", lambda task_id, worker_id: None)
    # Allowed types, with content checks.
    r = _upload(admin, "notes.md", b"# Quarterly plan\nShip the thing.")
    assert r.status_code == 201, r.text
    md = r.json()
    assert md["kind"] == "document" and md["text_chars"] > 10
    xl = _upload(admin, "sales.xlsx", _xlsx()).json()
    assert xl["kind"] == "document" and xl["text_chars"] > 0
    img = _upload(admin, "photo.png", _png()).json()
    assert img["kind"] == "image" and img["mime"] == "image/png"
    assert _upload(admin, "evil.exe", b"MZ\x90\x00").status_code == 422
    fake = _upload(admin, "fake.png", b"<script>alert(1)</script>")
    assert fake.status_code == 422 and "does not look like a PNG" in fake.text
    assert _upload(admin, "bin.txt", b"abc\x00def").status_code == 422
    assert _upload(admin, "empty.pdf", b"").status_code == 422
    assert _upload(admin, "../../etc/passwd.txt", b"x").json()["filename"] == "passwd.txt"

    # Before a task uses it, only the uploader can see an attachment.
    admin.post("/api/users", {"email": "att-viewer@acme.com", "name": "Vic Viewer", "role": "operator", "password": PASSWORD})
    other = login_as(server, "att-viewer@acme.com")
    assert other.get(f"/api/attachments/{img['id']}").status_code == 404
    got = admin.get(f"/api/attachments/{img['id']}?inline=true")
    assert got.status_code == 200 and got.headers["content-type"] == "image/png"
    assert got.headers["content-disposition"].startswith("inline") and "sandbox" in got.headers["content-security-policy"]
    doc = admin.get(f"/api/attachments/{md['id']}?inline=true")
    assert doc.headers["content-disposition"].startswith("attachment")  # documents always download

    # Submitting links them to the task; they appear on it and can't be reused or deleted.
    agent = make_agent(admin, offline_provider, "Attachment Reader")
    t = admin.post(f"/api/agents/{agent['id']}/tasks", {"instructions": "Summarise the attached files",
                                                         "attachment_ids": [md["id"], xl["id"], img["id"]]})
    assert t.status_code == 201, t.text
    assert [a["filename"] for a in t.json()["attachments"]] == ["notes.md", "sales.xlsx", "photo.png"]
    again = admin.post(f"/api/agents/{agent['id']}/tasks", {"instructions": "x", "attachment_ids": [md["id"]]})
    assert again.status_code == 422 and "already attached" in again.text
    assert admin.delete(f"/api/attachments/{md['id']}").status_code == 409
    assert other.get(f"/api/attachments/{img['id']}").status_code == 200  # now visible with the task
    stranger = _upload(other, "theirs.txt", b"hello").json()
    r = admin.post(f"/api/agents/{agent['id']}/tasks", {"instructions": "x", "attachment_ids": [stranger["id"]]})
    assert r.status_code == 422  # someone else's upload

    # Unsent uploads can be removed by their uploader, and are purged after a day.
    tmp = _upload(admin, "scratch.txt", b"temp").json()
    assert admin.delete(f"/api/attachments/{tmp['id']}").status_code == 200
    old = _upload(admin, "old.txt", b"old").json()
    with session_scope() as db:
        db.get(Attachment, old["id"]).created_at = utcnow() - timedelta(days=2)
    from sca.services.retention import apply_retention

    assert apply_retention()["unsent_attachments"] >= 1
    with session_scope() as db:
        assert db.get(Attachment, old["id"]) is None


def test_agent_receives_documents_and_images(admin, offline_provider, tmp_path):
    from sca.orchestrator.executor import _scan_artifacts, _stage_attachments, _with_attachments
    from sca.services import attachments

    agent = make_agent(admin, offline_provider, "Vision Reader")
    long_text = "Revenue grew. " * 3000  # longer than the inline limit
    with session_scope() as db:
        t, _ = submit_task(db, agent=db.get(Agent, agent["id"]), instructions="Summarise")
        t.scheduled_for = utcnow() + timedelta(hours=1)  # keep it out of the worker
        user = t.requested_by or "u"
        a1 = attachments.save(db, t.org_id, user, "report.docx", _docx(long_text))
        a2 = attachments.save(db, t.org_id, user, "chart.png", _png(3000, 2000))
        a3 = attachments.save(db, t.org_id, user, "report.docx", _docx("second copy"))
        attachments.link(db, t.org_id, user, [a1.id, a2.id, a3.id], t)
        task_id = t.id
        staged = _stage_attachments(db, task_id, tmp_path)

    assert [a["path"] for a in staged] == ["attachments/report.docx", "attachments/chart.png", "attachments/report (2).docx"]
    assert (tmp_path / "attachments/report.docx.txt").read_text().startswith("Revenue grew.")

    from openhands.sdk import Message

    vision = SimpleNamespace(attachments=staged, llm=SimpleNamespace(vision_is_active=lambda: True))
    msg = _with_attachments(vision, "Summarise")
    assert isinstance(msg, Message)
    text = msg.content[0].text
    assert "attachments/report.docx" in text and "full text in attachments/report.docx.txt" in text
    assert "<untrusted_data>" in text and "chart.png (image, shown to you below)" in text
    url = msg.content[1].image_urls[0]
    assert url.startswith("data:image/jpeg;base64,")  # 3000px image downscaled for the model

    blind = _with_attachments(SimpleNamespace(attachments=staged, llm=SimpleNamespace(vision_is_active=lambda: False)), "Summarise")
    assert isinstance(blind, str) and "your model cannot view images" in blind
    assert _with_attachments(SimpleNamespace(attachments=[], llm=None), "Hi") == "Hi"

    # Copied inputs are not reported as files the agent produced.
    (tmp_path / "summary.md").write_text("done")
    seen: list[str] = []
    import sca.orchestrator.executor as ex

    p = SimpleNamespace(project_dir=tmp_path, attachments=staged, org_id="o", task_id=task_id, agent_id=agent["id"], session_id="s")
    orig = ex.bus.emit
    ex.bus.emit = lambda *a, **k: seen.extend(a[5]["files"]) if len(a) > 5 else None
    try:
        _scan_artifacts(p, since=0)
    finally:
        ex.bus.emit = orig
    assert seen == ["summary.md"]
